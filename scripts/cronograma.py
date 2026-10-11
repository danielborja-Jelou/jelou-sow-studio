#!/usr/bin/env python3
"""Cronograma del SOW: Gantt en Excel editable + imagen/PDF para adjuntar.

Uso: cronograma.py <cronograma.json> [salida_base] [--sin-png] [--sin-pdf]
Genera:
  <base>.xlsx              Gantt editable (hoja Cronograma + hoja Entregables)
  <base>_gantt.svg         Gantt dibujado con los colores y la tipografía de assets/brand.json
  <base>_gantt.png         imagen (usa scripts/svg_to_png.js, el mismo motor resvg de los diagramas)
  <base>_gantt.pdf         adjunto para el correo (necesita Chrome o Edge; si no hay, se omite)
  <base>_entregables.json  hitos con fecha

Formato del JSON: assets/cronograma_ejemplo.json (lo deriva sow_a_paquete.py desde sow.json).
Las fechas se calculan en DÍAS HÁBILES (sin fines de semana ni `feriados`): cada tarea arranca el día
hábil siguiente al fin de sus dependencias (`depende_de`) o en `inicio` si no tiene.
`responsable` = "Jelou" | "Cliente" | "Ambos" (cambia el color). `hito: true` = entregable con fecha (0 días).
Sin dependencias de Mermaid ni de servicios web: todo se dibuja localmente.
"""
import datetime as dt
import json
import os
import shutil
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import brand  # noqa: E402

PAL = brand.paleta()
FUENTE, FUENTE_BOLD = brand.fuente()
# Jelou = turquesa de marca, Cliente = gris de marca, Ambos = turquesa oscuro, hito = oscuro de marca
COLOR = {"Jelou": PAL["teal"], "Cliente": PAL["gris"], "Ambos": PAL["teal_oscuro"]}
COLOR_HITO = PAL["oscuro"]


def habil(d, feriados):
    return d.weekday() < 5 and d not in feriados


def siguiente_habil(d, feriados):
    d += dt.timedelta(days=1)
    while not habil(d, feriados):
        d += dt.timedelta(days=1)
    return d


def calcular(cfg):
    fer = {dt.date.fromisoformat(f) for f in cfg.get("feriados", [])}
    inicio = dt.date.fromisoformat(cfg["inicio"])
    while not habil(inicio, fer):
        inicio += dt.timedelta(days=1)
    por_id, out = {}, []
    for t in cfg["tareas"]:
        deps = t.get("depende_de", [])
        faltan = [d for d in deps if d not in por_id]
        if faltan:
            sys.exit(f"Tarea {t['id']}: depende de {faltan}, que no existe o va después en la lista")
        ini = max((siguiente_habil(por_id[d]["fin"], fer) for d in deps), default=inicio)
        if t.get("inicio"):
            ini = max(ini, dt.date.fromisoformat(t["inicio"]))
        dias = 0 if t.get("hito") else int(t["dias"])
        fin, n = ini, 1
        while n < dias:
            fin = siguiente_habil(fin, fer)
            n += 1
        r = {**t, "inicio": ini, "fin": fin, "dias": dias, "responsable": t.get("responsable", "Jelou")}
        por_id[t["id"]] = r
        out.append(r)
    return out, fer


def dias_habiles_entre(f0, d, fer):
    n, x = 1, f0
    while x < d:
        x = siguiente_habil(x, fer)
        n += 1
    return n


def entregables(tareas, f0, fer):
    out = []
    for t in tareas:
        if not t.get("hito"):
            continue
        out.append({"entregable": t.get("entregable", t["tarea"]), "descripcion": t.get("descripcion", ""),
                    "dias_habiles": dias_habiles_entre(f0, t["inicio"], fer), "fecha": t["inicio"].isoformat()})
    return out


def lista_dias(tareas, fer):
    f0, f1 = min(t["inicio"] for t in tareas), max(t["fin"] for t in tareas)
    dias, d = [], f0
    while d <= f1:
        if habil(d, fer):
            dias.append(d)
        d += dt.timedelta(days=1)
    return dias


# ---------- Excel editable ----------

def excel(cfg, tareas, fer, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cronograma"
    thin = Side(style="thin", color=PAL["borde"])
    hdr = PatternFill("solid", fgColor=PAL["oscuro"])
    f0 = min(t["inicio"] for t in tareas)
    dias = lista_dias(tareas, fer)
    fijas = ["Fase", "Tarea", "Responsable", "Inicio", "Fin", "Días hábiles", "Depende de"]
    ws.append([f"Cronograma — {cfg['cliente']} · {cfg.get('proyecto', '')} · v{cfg.get('version', '1.0')}"])
    ws["A1"].font = Font(bold=True, size=14, color=PAL["oscuro"], name=FUENTE_BOLD)
    sem = [""] * len(fijas) + [f"Sem {((x - f0).days // 7) + 1}" if (i == 0 or x.weekday() == 0) else "" for i, x in enumerate(dias)]
    ws.append(sem)
    ws.append(fijas + [x.strftime("%d/%m") for x in dias])
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF", name=FUENTE_BOLD, size=9)
        c.fill = hdr
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, text_rotation=90 if c.column > len(fijas) else 0)
    for c in ws[2]:
        c.font = Font(bold=True, color=PAL["teal_oscuro"], name=FUENTE_BOLD, size=9)
    for t in tareas:
        ws.append([t["fase"], ("◆ " if t.get("hito") else "") + t["tarea"], t["responsable"], t["inicio"], t["fin"],
                   t["dias"], ", ".join(t.get("depende_de", []))])
        i = ws.max_row
        for c in ws[i][:len(fijas)]:
            c.font = Font(name=FUENTE, size=10, bold=bool(t.get("hito")), color=PAL["ink"])
            c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical="center", wrap_text=c.column == 2)
        ws.cell(i, 4).number_format = ws.cell(i, 5).number_format = "dd/mm/yyyy"
        color = COLOR_HITO if t.get("hito") else COLOR.get(t["responsable"], PAL["gris"])
        for j, x in enumerate(dias):
            c = ws.cell(i, len(fijas) + 1 + j)
            c.border = Border(left=thin, bottom=thin)
            if t["inicio"] <= x <= t["fin"]:
                c.fill = PatternFill("solid", fgColor=color)
                if t.get("hito"):
                    c.value = "◆"
                    c.font = Font(color="FFFFFF", bold=True)
                    c.alignment = Alignment(horizontal="center")
    for i, w in enumerate([24, 42, 12, 11, 11, 8, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for j in range(len(dias)):
        ws.column_dimensions[get_column_letter(len(fijas) + 1 + j)].width = 3.2
    ws.row_dimensions[3].height = 42
    ws.freeze_panes = ws.cell(4, 3)
    fila = ws.max_row + 2
    ws.cell(fila, 1, "Leyenda").font = Font(bold=True, name=FUENTE_BOLD)
    for k, (txt, col) in enumerate([("Jelou", COLOR["Jelou"]), ("Cliente", COLOR["Cliente"]), ("Ambos", COLOR["Ambos"]), ("◆ Hito / entregable", COLOR_HITO)], 1):
        ws.cell(fila + k, 1).fill = PatternFill("solid", fgColor=col)
        ws.cell(fila + k, 2, txt).font = Font(name=FUENTE, size=10)
    ws.cell(fila + 6, 1, "Fechas en días hábiles; quedan sujetas a la entrega de insumos del cliente.").font = Font(italic=True, size=9, color=PAL["secundario"])
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    e = wb.create_sheet("Entregables")
    e.append(["Entregable (hito)", "Descripción", "Días hábiles de implementación", "Fecha de entrega"])
    for c in e[1]:
        c.font = Font(bold=True, color="FFFFFF", name=FUENTE_BOLD)
        c.fill = hdr
    for h in entregables(tareas, f0, fer):
        e.append([h["entregable"], h["descripcion"], h["dias_habiles"], dt.date.fromisoformat(h["fecha"])])
        e.cell(e.max_row, 4).number_format = "dd/mm/yyyy"
    for i, w in enumerate([34, 50, 16, 16], 1):
        e.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)


# ---------- Gantt SVG (sin Mermaid) ----------

def esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def recorta(s, n):
    s = str(s)
    return s if len(s) <= n else s[: n - 1] + "…"


def gantt_svg(cfg, tareas, fer):
    dias = lista_dias(tareas, fer)
    f0 = dias[0]
    idx = {d: i for i, d in enumerate(dias)}
    margen, col_fase, col_tarea, ancho_dia, alto_fila, alto_cab = 24, 210, 310, 24, 26, 58
    x0 = margen + col_fase + col_tarea
    ancho = x0 + len(dias) * ancho_dia + margen
    filas = len(tareas)
    y0 = margen + 44 + alto_cab
    alto = y0 + filas * alto_fila + 70
    o = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{ancho}" height="{alto}" viewBox="0 0 {ancho} {alto}" font-family="{FUENTE}, sans-serif">',
         f'<rect width="{ancho}" height="{alto}" fill="#ffffff"/>',
         f'<text x="{margen}" y="{margen + 18}" font-size="17" font-family="{FUENTE_BOLD}" fill="#{PAL["ink"]}">{esc("Cronograma — " + cfg["cliente"] + " · v" + str(cfg.get("version", "1.0")))}</text>',
         f'<text x="{margen}" y="{margen + 38}" font-size="11" fill="#{PAL["secundario"]}">{esc(cfg.get("proyecto", ""))}</text>']
    # cabecera: semanas y días
    ytop = margen + 44
    o.append(f'<rect x="{margen}" y="{ytop}" width="{ancho - 2 * margen}" height="{alto_cab}" fill="#{PAL["oscuro"]}"/>')
    for txt, x in (("Fase", margen + 8), ("Tarea", margen + col_fase + 8)):
        o.append(f'<text x="{x}" y="{ytop + 34}" font-size="11" font-family="{FUENTE_BOLD}" fill="#ffffff">{txt}</text>')
    for i, d in enumerate(dias):
        x = x0 + i * ancho_dia
        if (i == 0 or d.weekday() == 0) and x + 44 <= ancho - margen:
            o.append(f'<text x="{x + 2}" y="{ytop + 20}" font-size="10" font-family="{FUENTE_BOLD}" fill="#ffffff">Sem {((d - f0).days // 7) + 1}</text>')
        o.append(f'<text x="{x + ancho_dia / 2}" y="{ytop + 46}" font-size="8" text-anchor="middle" fill="#ffffff">{d.strftime("%d/%m")}</text>')
    # filas
    fase_prev = None
    for k, t in enumerate(tareas):
        y = y0 + k * alto_fila
        if k % 2 == 0:
            o.append(f'<rect x="{margen}" y="{y}" width="{ancho - 2 * margen}" height="{alto_fila}" fill="#{PAL["cabecera"]}"/>')
        if t["fase"] != fase_prev:
            o.append(f'<text x="{margen + 8}" y="{y + 17}" font-size="10" font-family="{FUENTE_BOLD}" fill="#{PAL["teal_oscuro"]}">{esc(recorta(t["fase"], 34))}</text>')
            fase_prev = t["fase"]
        nombre = t["tarea"] + ("" if t["responsable"] == "Jelou" or t.get("hito") else f" ({t['responsable']})")
        o.append(f'<text x="{margen + col_fase + 8}" y="{y + 17}" font-size="10" fill="#{PAL["ink"]}">{esc(recorta(nombre, 50))}</text>')
        cx = x0 + idx[t["inicio"]] * ancho_dia
        if t.get("hito"):
            c = COLOR_HITO
            mx, my = cx + ancho_dia / 2, y + alto_fila / 2
            o.append(f'<polygon points="{mx},{my - 8} {mx + 8},{my} {mx},{my + 8} {mx - 8},{my}" fill="#{c}"/>')
        else:
            c = COLOR.get(t["responsable"], PAL["gris"])
            w = (idx[t["fin"]] - idx[t["inicio"]] + 1) * ancho_dia
            o.append(f'<rect x="{cx + 1}" y="{y + 5}" width="{w - 2}" height="{alto_fila - 10}" rx="4" fill="#{c}"/>')
    # rejilla vertical (semanas)
    for i, d in enumerate(dias):
        if d.weekday() == 0:
            o.append(f'<line x1="{x0 + i * ancho_dia}" y1="{y0}" x2="{x0 + i * ancho_dia}" y2="{y0 + filas * alto_fila}" stroke="#{PAL["borde"]}" stroke-width="1"/>')
    # leyenda
    yl = y0 + filas * alto_fila + 24
    x = margen
    for txt, col in (("Jelou", COLOR["Jelou"]), ("Cliente", COLOR["Cliente"]), ("Ambos", COLOR["Ambos"]), ("Hito / entregable", COLOR_HITO)):
        o.append(f'<rect x="{x}" y="{yl - 10}" width="16" height="12" rx="3" fill="#{col}"/>')
        o.append(f'<text x="{x + 22}" y="{yl}" font-size="10" fill="#{PAL["ink"]}">{txt}</text>')
        x += 40 + len(txt) * 7
    o.append(f'<text x="{margen}" y="{yl + 24}" font-size="9" fill="#{PAL["secundario"]}">Fechas en días hábiles; quedan sujetas a la entrega de insumos del cliente.</text>')
    o.append("</svg>")
    return "\n".join(o)


def a_png(svg_path, png_path):
    js = os.path.join(os.path.dirname(os.path.abspath(__file__)), "svg_to_png.js")
    try:
        r = subprocess.run(["node", js, svg_path, png_path, "1600"], capture_output=True, text=True, timeout=120)
    except (OSError, subprocess.TimeoutExpired) as ex:
        print(f"PNG omitido (node no disponible): {ex}")
        return False
    if r.returncode != 0:
        print("PNG omitido: ejecuta `npm install` en scripts/ (necesita @resvg/resvg-js).\n" + r.stderr.strip()[:300])
        return False
    return True


def a_pdf(svg_path, pdf_path):
    """PDF de una página con el tamaño del SVG, vía Chrome/Edge headless."""
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from propuesta import navegador  # noqa: E402
    exe = navegador()
    if not exe:
        print("PDF omitido: no hay Chrome ni Edge para imprimir.")
        return False
    import re
    svg = open(svg_path, encoding="utf-8").read()
    m = re.search(r'width="(\d+)" height="(\d+)"', svg)
    w, h = (int(m.group(1)), int(m.group(2))) if m else (1600, 900)
    # HTML temporal en la carpeta temporal del sistema (ruta corta: en Windows una ruta larga hace fallar a Chrome)
    import tempfile
    tmp = tempfile.mkdtemp(prefix="gantt_")
    html_path = os.path.join(tmp, "g.html")
    with open(html_path, "w", encoding="utf-8") as fh:
        fh.write(f"<!doctype html><meta charset=utf-8><style>@page{{size:{w}px {h}px;margin:0}}html,body{{margin:0}}svg{{display:block}}</style>{svg}")
    ok = False
    try:
        r = subprocess.run([exe, "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                            f"--print-to-pdf={os.path.join(tmp, 'g.pdf')}", "file:///" + os.path.abspath(html_path).replace("\\", "/")],
                           capture_output=True, text=True, timeout=120)
        ok = r.returncode == 0 and os.path.exists(os.path.join(tmp, "g.pdf"))
        if ok:
            shutil.move(os.path.join(tmp, "g.pdf"), pdf_path)
    except (OSError, subprocess.TimeoutExpired):
        ok = False
    shutil.rmtree(tmp, ignore_errors=True)
    return ok


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit(__doc__)
    cfg = json.load(open(args[0], encoding="utf-8"))
    base = args[1] if len(args) > 1 else os.path.splitext(args[0])[0]
    os.makedirs(os.path.dirname(os.path.abspath(base)), exist_ok=True)
    tareas, fer = calcular(cfg)
    excel(cfg, tareas, fer, base + ".xlsx")
    with open(base + "_gantt.svg", "w", encoding="utf-8") as f:
        f.write(gantt_svg(cfg, tareas, fer))
    f0 = min(t["inicio"] for t in tareas)
    json.dump(entregables(tareas, f0, fer), open(base + "_entregables.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    salidas = [".xlsx", "_gantt.svg", "_entregables.json"]
    if "--sin-png" not in sys.argv and a_png(base + "_gantt.svg", base + "_gantt.png"):
        salidas.append("_gantt.png")
    if "--sin-pdf" not in sys.argv and a_pdf(base + "_gantt.svg", base + "_gantt.pdf"):
        salidas.append("_gantt.pdf")
    for ext in salidas:
        print(base + ext)
    print(f"Inicio {f0:%d/%m/%Y} - fin {max(t['fin'] for t in tareas):%d/%m/%Y}")


if __name__ == "__main__":
    main()
